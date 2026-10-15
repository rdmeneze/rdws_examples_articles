#!/usr/bin/env python3
"""Junta os .tsv de scripts/matriz.sh em uma planilha: recursos nas linhas, toolchains nas
colunas, com uma aba de legenda. Precisa do openpyxl (pip install openpyxl, ou o pacote
python3-openpyxl da distribuição).

Uso: scripts/tabela.py resultados/matriz.xlsx resultados/*.tsv
"""
import pathlib
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TESTES = pathlib.Path(__file__).resolve().parent.parent / "testes"

CORES = {
    "ok": "C6EFCE",
    "não": "FFC7CE",
    "só gnu": "FFEB9C",
    "só declara": "FFEB9C",
    "falha": "FFC7CE",
    "n/t": "E7E6E6",
}
SIGNIFICADO = [
    ("ok", "Funciona: compila, linka ou roda, conforme o teste pede."),
    ("não", "Não compila com -std=c23 (ou c2x), -pedantic-errors e -Werror."),
    ("só gnu", "Recurso de biblioteca que existe e linka, mas a libc só o declara com -std=gnu23."),
    ("só declara", "Compila, mas não linka: o cabeçalho declara a função e a biblioteca não a tem."),
    ("falha", "Roda, mas com a saída errada."),
    ("n/t", "Não testado: teste de execução em um alvo cruzado, onde o binário não roda."),
]
METADADOS = {"_compilador": "Compilador", "_libc": "Biblioteca C", "_std": "Padrão (__STDC_VERSION__)"}

FONTE = Font(name="Arial", size=10)
NEGRITO = Font(name="Arial", size=10, bold=True)
ITALICO = Font(name="Arial", size=10, italic=True)
TITULO = Font(name="Arial", size=12, bold=True)
FINA = Side(style="thin", color="BFBFBF")
BORDA = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)
CINZA = PatternFill("solid", fgColor="F2F2F2")
CABECALHO = PatternFill("solid", fgColor="D9E1F2")


def descricao(recurso):
    """Lê tipo, verificação e descrição da primeira linha do arquivo de teste."""
    arquivo = TESTES / f"{recurso}.c"
    if not arquivo.exists():
        return recurso, "", ""
    cab = arquivo.read_text(encoding="utf-8").splitlines()[0]
    campo = lambda nome: (re.search(rf"{nome}: ([^;*]+)", cab) or [None, ""])[1].strip()
    return campo("recurso") or recurso, campo("tipo"), campo("verifica")


def le_resultados(caminhos):
    colunas, recursos, dados = [], [], {}
    for caminho in caminhos:
        for linha in pathlib.Path(caminho).read_text(encoding="utf-8").splitlines():
            nome, recurso, res = linha.split("\t")
            if nome not in colunas:
                colunas.append(nome)
            if recurso not in recursos:
                recursos.append(recurso)
            dados[(recurso, nome)] = res
    return colunas, recursos, dados


def celula(ws, linha, coluna, valor, fonte=FONTE, preenchimento=None, alinhamento=None):
    c = ws.cell(row=linha, column=coluna, value=valor)
    c.font = fonte
    c.border = BORDA
    if preenchimento:
        c.fill = preenchimento
    c.alignment = alinhamento or Alignment(vertical="center", wrap_text=True)
    return c


def aba_matriz(ws, colunas, recursos, dados):
    ws.title = "Matriz"
    n_fixas = 3
    ultima = n_fixas + len(colunas)
    ws.cell(row=1, column=1, value="C23 em sistemas embarcados: recurso × toolchain").font = TITULO
    ws.cell(row=2, column=1,
            value="Gerado por scripts/rodar_tudo.sh; legenda e método na aba Legenda.").font = ITALICO

    linha = 4
    for j, texto in enumerate(["Recurso", "Tipo", "Verifica"] + colunas, start=1):
        celula(ws, linha, j, texto, NEGRITO, CABECALHO, Alignment(horizontal="center", vertical="center", wrap_text=True))

    metas = [r for r in recursos if r.startswith("_")]
    testes = [r for r in recursos if not r.startswith("_")]
    for r in metas:
        linha += 1
        celula(ws, linha, 1, METADADOS.get(r, r), ITALICO, CINZA)
        celula(ws, linha, 2, "", ITALICO, CINZA)
        celula(ws, linha, 3, "", ITALICO, CINZA)
        for j, c in enumerate(colunas, start=n_fixas + 1):
            celula(ws, linha, j, dados.get((r, c), ""), ITALICO, CINZA, Alignment(horizontal="center", vertical="center", wrap_text=True))

    primeira_teste = linha + 1
    for r in testes:
        linha += 1
        nome, tipo, verifica = descricao(r)
        celula(ws, linha, 1, nome)
        celula(ws, linha, 2, tipo)
        celula(ws, linha, 3, verifica)
        for j, c in enumerate(colunas, start=n_fixas + 1):
            res = dados.get((r, c), "")
            cor = PatternFill("solid", fgColor=CORES[res]) if res in CORES else None
            celula(ws, linha, j, res, FONTE, cor, Alignment(horizontal="center", vertical="center"))
    ultima_teste = linha

    linha += 1
    celula(ws, linha, 1, f"Recursos com \"ok\" (de {len(testes)})", NEGRITO, CABECALHO)
    celula(ws, linha, 2, "", NEGRITO, CABECALHO)
    celula(ws, linha, 3, "", NEGRITO, CABECALHO)
    for j in range(n_fixas + 1, ultima + 1):
        col = get_column_letter(j)
        celula(ws, linha, j, f'=COUNTIF({col}{primeira_teste}:{col}{ultima_teste},"ok")', NEGRITO, CABECALHO,
               Alignment(horizontal="center", vertical="center"))

    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 9
    for j in range(n_fixas + 1, ultima + 1):
        ws.column_dimensions[get_column_letter(j)].width = 15
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = ws.cell(row=primeira_teste, column=n_fixas + 1)
    # Impressão: paisagem, todas as colunas em uma página de largura.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def aba_legenda(wb):
    ws = wb.create_sheet("Legenda")
    ws.cell(row=1, column=1, value="Legenda").font = TITULO
    linha = 3
    for res, texto in SIGNIFICADO:
        celula(ws, linha, 1, res, NEGRITO, PatternFill("solid", fgColor=CORES[res]), Alignment(horizontal="center", vertical="center"))
        celula(ws, linha, 2, texto)
        linha += 1

    linha += 1
    ws.cell(row=linha, column=1, value="Método").font = TITULO
    notas = [
        "Um arquivo por recurso em testes/; a primeira linha diz o tipo e o que verificar (compila, linka ou roda).",
        "Compilação com -std=c23 (ou c2x), -pedantic-errors e -Werror: extensão do compilador não conta como suporte.",
        "Recurso de biblioteca que falha no modo estrito é testado de novo com -std=gnu23 (\"só gnu\").",
        "Host: imagens Docker gcc:13 a 16, silkeh/clang:18 a 22 e fedora:44 (glibc 2.43).",
        "ARM: Arm GNU Toolchain 13.3, 14.3 e 15.2 (newlib) e Zephyr SDK 1.0.1 (picolibc), para Cortex-M4, sem execução.",
        "Nas colunas de host, a libc é a da imagem Docker: as imagens do Clang 18 a 20 trazem a glibc 2.36.",
        "Código e scripts: https://github.com/rdmeneze/rdws_examples_articles/tree/main/modern-c/02-c23-em-sistemas-embarcados",
    ]
    for texto in notas:
        linha += 1
        ws.cell(row=linha, column=1, value="•").font = FONTE
        c = ws.cell(row=linha, column=2, value=texto)
        c.font = FONTE
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 100


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    saida, entradas = sys.argv[1], sys.argv[2:]
    colunas, recursos, dados = le_resultados(entradas)
    wb = Workbook()
    aba_matriz(wb.active, colunas, recursos, dados)
    aba_legenda(wb)
    wb.save(saida)
    print(f"tabela: {saida}")


if __name__ == "__main__":
    main()
